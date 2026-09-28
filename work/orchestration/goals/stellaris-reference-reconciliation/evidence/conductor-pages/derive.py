"""Source reconciliation arithmetic; no production model imports or tuning."""
import json
from pathlib import Path
from openpyxl import load_workbook
P=Path(__file__).resolve().parent
I=15.4e6
j=I/.36**2/1e6
nref=50000/j*.09/.336
ic=lambda b:300*(b/20)**(-.6)
r={"reference_density_A_mm2":j,"entering_turns":I/50000,"reference_tapes":nref,"source_cell_tapes_at_assumed_56um":6/.056,"source_turn_current_from_rounded_amp_turns_A":I/324,"load_per_tape_A":50000/nref,"set_tape_length_m":136.56*.09/(.006*.000056),"set_conductor_length_m":48*I*(80.4/92.4)*25/50000,"source_ungraded_length_m":8*sum([807,847,721,717,636,567])*1000,"source_graded_length_m":8*sum([167,161,146,134,131,118])*1000,"source_total_turns":8*sum([324,324,289,289,256,225])}
r['set_tapes']=r['set_tape_length_m']/r['set_conductor_length_m']
r['field_attribution']=[dict(B_T=b,Ic_tape_A=ic(b),Ic_cable_A=nref*ic(b),operating_fraction=50000/(nref*ic(b)),allowable_margin_A=.8*nref*ic(b)-50000) for b in [24.9,24.7,24.6,24.59]]
r['rounded_source_cell_scalar_fraction']=47600/((6/.056)*ic(24.6))
r['source_minimum_local_Ic_from_rounded_turn_current_A']=47600/.605
r['source_minimum_local_Je_from_printed_stack_J_A_mm2']=1323/.605
r['unresolved_ratio_not_an_orientation_factor']=r['field_attribution'][0]['operating_fraction']/.605
w=load_workbook(P/'wimbush-20K.xlsx',data_only=True)
r['dataset_sheet_names']=w.sheetnames
r['dataset_angular_examples']={}
for name in ['5T','8T']:
 rows=list(w[name].values)[1:]
 take=lambda row:dict(Ic_per_cm_A=row[3],Hall_angle_deg=row[9],set_angle_deg=row[12])
 lo=min(rows,key=lambda x:x[3]);hi=max(rows,key=lambda x:x[3])
 r['dataset_angular_examples'][name]=dict(minimum=take(lo),maximum=take(hi),ratio=hi[3]/lo[3])
(P/'arithmetic.json').write_text(json.dumps(r,indent=2)+'\n')
print(json.dumps(r,indent=2))
